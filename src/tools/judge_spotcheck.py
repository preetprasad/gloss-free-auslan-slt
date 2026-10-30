"""t3-errs6 Part 3: validate the adequacy judge against a human (Preet) on 50 S6 test outputs.

make  : draw 10 clips from each of the judge's grades 0-4 (seed 0), shuffle, and write a blind grading sheet
        (results/semantic_regrade_10288/judge_spotcheck_SHEET.xlsx) plus a separate key the grader must not open
        (judge_spotcheck_KEY.json).
score : read the filled sheet and report Spearman rho, exact agreement, within-one agreement and the confusion
        matrix against the judge. Pre-registered in results/PREREG_WAVE4_2026-09-22.md section 7.
"""
import json, random, sys
from collections import Counter
from pathlib import Path

D = Path(__file__).resolve().parents[2] / "results" / "semantic_regrade_10288"
SHEET, KEY = D / "judge_spotcheck_SHEET.xlsx", D / "judge_spotcheck_KEY.json"
RUBRIC = ["Grade how much of the REFERENCE meaning the SYSTEM OUTPUT conveys. Judge meaning only: ignore grammar, "
          "word order and exact wording.",
          "0 = unrelated or empty", "1 = shares a topic word only", "2 = partial meaning",
          "3 = most of the meaning, minor errors", "4 = fully adequate (a correct paraphrase counts as 4)",
          "Grade each row on its own. Do not look at judge_spotcheck_KEY.json until every row is graded."]


def make():
    import csv, openpyxl
    rows = list(csv.DictReader(open(D / "s6.graded_hf.csv")))
    rng, picked = random.Random(0), []
    for g in range(5):
        pool = [r for r in rows if int(r["adequacy"]) == g]
        picked += rng.sample(pool, 10)
    rng.shuffle(picked)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "grade these"
    ws.append(["item", "reference (what was signed)", "system output", "your grade 0-4", "notes"])
    for n, r in enumerate(picked, 1):
        ws.append([n, r["ref"], r["hyp"], None, None])
    for col, w in zip("ABCDE", (6, 55, 55, 14, 30)):
        ws.column_dimensions[col].width = w
    rb = wb.create_sheet("rubric")
    for line in RUBRIC:
        rb.append([line])
    rb.column_dimensions["A"].width = 110
    wb.save(SHEET)
    json.dump([{"item": n, "idx": int(r["idx"]), "judge": int(r["adequacy"])} for n, r in enumerate(picked, 1)],
              open(KEY, "w"), indent=1)
    print(f"wrote {SHEET} (50 rows, 10 per judge grade, shuffled) and the key {KEY}")


def spearman(a, b):
    def ranks(x):
        order = sorted(range(len(x)), key=lambda i: x[i])
        r, i = [0.0] * len(x), 0
        while i < len(x):
            j = i
            while j + 1 < len(x) and x[order[j + 1]] == x[order[i]]:
                j += 1
            for k in range(i, j + 1):
                r[order[k]] = (i + j) / 2 + 1
            i = j + 1
        return r
    ra, rb = ranks(a), ranks(b)
    ma, mb = sum(ra) / len(ra), sum(rb) / len(rb)
    cov = sum((x - ma) * (y - mb) for x, y in zip(ra, rb))
    return cov / (sum((x - ma) ** 2 for x in ra) * sum((y - mb) ** 2 for y in rb)) ** 0.5


def score():
    import openpyxl
    key = {k["item"]: k["judge"] for k in json.load(open(KEY))}
    ws = openpyxl.load_workbook(SHEET)["grade these"]
    human = {int(r[0]): r[3] for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None}
    missing = [i for i in key if human.get(i) not in (0, 1, 2, 3, 4)]
    assert not missing, f"items not graded 0-4: {missing}"
    h, j = [int(human[i]) for i in sorted(key)], [key[i] for i in sorted(key)]
    exact = sum(a == b for a, b in zip(h, j)) / len(h)
    within1 = sum(abs(a - b) <= 1 for a, b in zip(h, j)) / len(h)
    print(f"n {len(h)} | Spearman rho {spearman(h, j):.3f} | exact {100 * exact:.0f}% | within one {100 * within1:.0f}%")
    print("confusion (rows = judge 0-4, cols = human 0-4):")
    c = Counter(zip(j, h))
    for g in range(5):
        print(g, [c[(g, x)] for x in range(5)])


def selftest():
    assert abs(spearman([1, 2, 3], [1, 2, 3]) - 1) < 1e-12 and abs(spearman([1, 2, 3], [3, 2, 1]) + 1) < 1e-12
    assert abs(spearman([1, 1, 2, 2], [1, 1, 2, 2]) - 1) < 1e-12
    print("selftest passed")


if __name__ == "__main__":
    {"make": make, "score": score, "selftest": selftest}[sys.argv[1]]()
