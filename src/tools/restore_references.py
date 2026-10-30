#!/usr/bin/env python3
"""Put the Auslan-Daily reference sentences back into the saved outputs.

No sentence of the Auslan-Daily dataset is stored in this repository. In the saved outputs under results/,
each reference sentence is replaced by its SHA-1 digest ("sha1:<40 hex digits>"). This tool reads the
dataset's own translation spreadsheet, which `make data` fetches from the dataset's official Google Drive
folder, puts the sentences back, and writes the text manifests that some checks read.

    python tools/restore_references.py              # restore; verified against results/SHA256SUMS
    python tools/restore_references.py --self_test

Every restored file must have the SHA-256 recorded in results/SHA256SUMS, so a wrong or changed spreadsheet
is detected rather than silently used. The inverse (`strip_text`) is applied when the repository is assembled.

Dataset: X. Shen, S. Yuan, H. Sheng, H. Du and X. Yu, "Auslan-Daily: Australian Sign Language Translation for
Daily Communication and News", NeurIPS 2023 Datasets and Benchmarks. See ATTRIBUTION.md.
"""
import argparse, hashlib, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
XLSX = os.path.join(REPO, "data", "auslan-daily", "translation", "Auslan-Daily_Communication.xlsx")
RESULTS = os.path.join(REPO, "results")
MARKER = os.path.join(RESULTS, ".references_restored")
LITERAL = r'"(?:[^"\\]|\\.)*"'                       # one JSON string literal
PLACEHOLDER = re.compile(r'"sha1:([0-9a-f]{40})"')
SPLITS = {"train": "train", "dev": "val", "val": "val", "valid": "val", "validation": "val", "test": "test"}


def digest(text):
    return hashlib.sha1(text.encode("utf-8")).hexdigest()


def strip_text(text, name):
    """Replace every reference sentence in one saved-output file by its digest. Hypotheses are left alone."""
    hide = lambda lit: '"sha1:%s"' % digest(json.loads(lit))
    if name.endswith(".jsonl"):                      # one object per line with a "ref" field
        return re.sub(r'("ref": )(' + LITERAL + ")", lambda m: m.group(1) + hide(m.group(2)), text)
    if name.endswith(".json"):                       # a "refs" list
        return re.sub(r'("refs": \[)(.*?)(\])',
                      lambda m: m.group(1) + re.sub(LITERAL, lambda q: hide(q.group(0)), m.group(2)) + m.group(3),
                      text, flags=re.S)
    return text


def restore_text(text, table, ensure_ascii=True):
    missing = set()

    def put(m):
        if m.group(1) not in table:
            missing.add(m.group(1)); return m.group(0)
        return json.dumps(table[m.group(1)], ensure_ascii=ensure_ascii)
    return PLACEHOLDER.sub(put, text), missing


def sentences(xlsx):
    """[(clip name, subtitle, split)] in spreadsheet order, with the text exactly as tools/build_manifest.py writes it."""
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, read_only=True, data_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(rows)]
    ci, ti, si = head.index("Video_Clip_Name"), head.index("Subtitle"), head.index("Split")
    out = []
    for r in rows:
        if r[ti] is None:
            continue
        text = str(r[ti]).strip()
        if not text or text.lower() == "nan":
            continue
        out.append((str(r[ci]).strip(), text, SPLITS.get(str(r[si]).strip().lower(), str(r[si]).strip().lower())))
    return out


_TABLE = None


def sentence(sha1, xlsx=XLSX):
    """The Auslan-Daily subtitle that has this SHA-1 digest, read from the dataset's spreadsheet (`make data`).

    Scripts that need a particular sentence of the dataset (an example in a figure, a constant baseline) name it
    by digest and obtain the text here, so that the sentence itself is not stored in the repository."""
    global _TABLE
    if _TABLE is None:
        if not os.path.exists(xlsx):
            sys.exit(f"FATAL: {xlsx} not found. Run `make data` first (see ATTRIBUTION.md).")
        _TABLE = {digest(t): t for _, t, _ in sentences(xlsx)}
    return _TABLE[sha1]


def read_sums(path):
    return dict(reversed(line.rstrip("\n").split("  ", 1)) for line in open(path, encoding="utf-8") if line.strip())


def restore(xlsx, results=RESULTS, manifest_dir=os.path.join(REPO, "comm_run", "manifests")):
    if not os.path.exists(xlsx):
        sys.exit(f"FATAL: {xlsx} not found. Run `make data`, or download Auslan-Daily_Communication.xlsx from the "
                 "official Auslan-Daily folder into that location (see ATTRIBUTION.md).")
    rows = sentences(xlsx)
    table = {digest(t): t for _, t, _ in rows}
    sums = read_sums(os.path.join(results, "SHA256SUMS"))
    restored = already = 0
    for rel, want in sorted(sums.items()):
        path = os.path.join(results, rel)
        raw = open(path, "rb").read()
        if hashlib.sha256(raw).hexdigest() == want:
            already += 1; continue
        for ensure_ascii in (True, False):
            out, missing = restore_text(raw.decode("utf-8"), table, ensure_ascii)
            if missing:
                sys.exit(f"FATAL: {len(missing)} reference sentences of {rel} are not in {os.path.basename(xlsx)}; "
                         "is it the Communication translation spreadsheet of Auslan-Daily?")
            if hashlib.sha256(out.encode("utf-8")).hexdigest() == want:
                break
        else:
            sys.exit(f"FATAL: {rel} does not restore to its recorded SHA-256")
        open(path, "wb").write(out.encode("utf-8")); restored += 1
    os.makedirs(manifest_dir, exist_ok=True)
    written = []
    for split in ("train", "val", "test"):
        path = os.path.join(manifest_dir, split + ".tsv")
        if os.path.exists(path):                      # never overwrite a manifest that points at real pose files
            continue
        with open(path, "w", encoding="utf-8") as fh:
            fh.write("\n".join(f"dwpose_clips/{c}.npz\t{t}" for c, t, s in rows if s == split) + "\n")
        written.append(split)
    open(MARKER if results == RESULTS else os.path.join(results, ".references_restored"), "w").write("restored\n")
    print(f"{len(rows)} subtitles read; {restored} files restored, {already} already complete, all match results/SHA256SUMS; "
          f"text manifests written: {', '.join(written) or 'none (already present)'}")


def self_test():
    line = '{"idx": 0, "ref": "it is \\"fine\\" .", "hyp_beam": "it is fine ."}\n'
    doc = '{\n  "refs": [\n    "hello .",\n    "yes ."\n  ],\n  "beam_hyps": ["hello ."]\n}'
    table = {digest(t): t for t in ('it is "fine" .', "hello .", "yes .")}
    for name, text in (("a.jsonl", line), ("b.json", doc)):
        hidden = strip_text(text, name)
        assert "sha1:" in hidden and hidden != text
        assert '"hyp_beam": "it is fine ."' in hidden or '"beam_hyps": ["hello ."]' in hidden   # hypotheses untouched
        back, missing = restore_text(hidden, table)
        assert back == text and not missing, name
    assert restore_text('"sha1:%s"' % ("0" * 40), table)[1] == {"0" * 40}                        # unknown digest is reported
    print("restore_references self-test passed")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", default=XLSX)
    ap.add_argument("--self_test", action="store_true")
    a = ap.parse_args()
    self_test() if a.self_test else restore(a.xlsx)


if __name__ == "__main__":
    main()
